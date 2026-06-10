"""Shared openpyxl styling helpers for all generated workbooks."""

from __future__ import annotations

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# Colour palette
DARK_BLUE = "1F4E79"
MID_BLUE = "2E75B6"
LIGHT_BLUE = "D6E4F0"
WHITE = "FFFFFF"
PASS_GREEN = "C6EFCE"
WARN_AMBER = "FFEB9C"
FAIL_RED = "FFC7CE"
GREY_TEXT = "595959"
LIGHT_GREY = "F2F2F2"


def style_header_row(ws, row: int = 1, fill_colour: str = DARK_BLUE) -> None:
    fill = PatternFill(fill_type="solid", fgColor=fill_colour)
    font = Font(bold=True, color=WHITE, name="Calibri", size=11)
    align = Alignment(horizontal="center", vertical="center", wrap_text=False)
    for cell in ws[row]:
        cell.fill = fill
        cell.font = font
        cell.alignment = align


def style_section_header(ws, row: int, fill_colour: str = MID_BLUE) -> None:
    fill = PatternFill(fill_type="solid", fgColor=fill_colour)
    font = Font(bold=True, color=WHITE, name="Calibri", size=10)
    for cell in ws[row]:
        cell.fill = fill
        cell.font = font


def apply_alternate_rows(ws, start_row: int = 2, fill_colour: str = LIGHT_BLUE) -> None:
    alt_fill = PatternFill(fill_type="solid", fgColor=fill_colour)
    for i, row in enumerate(ws.iter_rows(min_row=start_row), start=start_row):
        if i % 2 == 0:
            for cell in row:
                if cell.fill.fill_type in ("none", None) or cell.fill.fgColor.rgb in (
                    "00000000", "FFFFFFFF"
                ):
                    cell.fill = alt_fill


def auto_column_width(ws, min_width: int = 8, max_width: int = 60) -> None:
    for col_cells in ws.columns:
        col_letter = get_column_letter(col_cells[0].column)
        best = max(
            (len(str(c.value)) for c in col_cells if c.value is not None),
            default=0,
        )
        ws.column_dimensions[col_letter].width = max(min_width, min(best + 4, max_width))


def freeze_and_filter(ws, freeze_cell: str = "A2") -> None:
    ws.freeze_panes = freeze_cell
    if ws.dimensions and ws.dimensions != "A1:A1":
        ws.auto_filter.ref = ws.dimensions


def colour_status_cells(ws, status_col_letter: str, start_row: int = 2) -> None:
    """Apply traffic-light fills to a STATUS column (PASS/WARN/FAIL)."""
    colour_map = {
        "PASS": PatternFill(fill_type="solid", fgColor=PASS_GREEN),
        "WARN": PatternFill(fill_type="solid", fgColor=WARN_AMBER),
        "FAIL": PatternFill(fill_type="solid", fgColor=FAIL_RED),
    }
    for row in ws.iter_rows(min_row=start_row):
        for cell in row:
            if cell.column_letter == status_col_letter and cell.value in colour_map:
                cell.fill = colour_map[cell.value]


def apply_full_style(ws, freeze_cell: str = "A2", status_col_letter: str | None = None) -> None:
    """Apply the full house style to a worksheet."""
    style_header_row(ws)
    auto_column_width(ws)
    freeze_and_filter(ws, freeze_cell)
    apply_alternate_rows(ws)
    if status_col_letter:
        colour_status_cells(ws, status_col_letter)
