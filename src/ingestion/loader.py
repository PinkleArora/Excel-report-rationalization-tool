"""Load Excel and CSV workbooks into a structured in-memory representation."""

from __future__ import annotations

import logging
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import pandas as pd

logger = logging.getLogger(__name__)

SUPPORTED_EXTENSIONS = {".xlsx", ".xlsm", ".xls", ".csv"}


@dataclass
class WorkbookBundle:
    """All sheets parsed from a single file, keyed by sheet name."""

    source_path: Path
    sheets: dict[str, pd.DataFrame] = field(default_factory=dict)
    # openpyxl Workbook object retained for formula / style inspection (None for CSV/xls)
    _raw_wb: object = field(default=None, repr=False, compare=False)

    @property
    def sheet_names(self) -> list[str]:
        return list(self.sheets.keys())

    @property
    def file_name(self) -> str:
        return self.source_path.name


def detect_header_row(raw_ws: Any, df_raw: pd.DataFrame, max_scan: int = 30) -> int:
    """Return the 0-based row index of the actual table header in a sheet.

    Detection order (first successful strategy wins):

    1. **Excel Table structure** — if the worksheet contains an Excel-formatted
       table (Insert → Table in Excel), its first row is the authoritative header.
    2. **Row scoring** — each row in the first *max_scan* rows is scored on:
       - text ratio (header labels are strings, not numbers)
       - uniqueness (column names should all be different)
       - density (most columns should be filled)
       - bonus when the immediately following row has numeric values

    Rows that are too sparse (< 25 % of columns filled) are skipped as title /
    banner / blank rows.

    Returns 0 when the best candidate is already on row 1.
    """
    # ── Strategy 1: Excel Table ───────────────────────────────────────────────
    if raw_ws is not None:
        try:
            tables = getattr(raw_ws, "tables", {})
            if tables:
                for tbl in (tables.values() if hasattr(tables, "values") else tables):
                    ref = getattr(tbl, "ref", "") or ""
                    m = re.match(r"[A-Za-z]+(\d+)", ref)
                    if m:
                        return int(m.group(1)) - 1  # 0-based
        except Exception:
            pass

    # ── Strategy 2: Row scoring ───────────────────────────────────────────────
    if df_raw.empty:
        return 0

    n_cols = len(df_raw.columns)
    scan_limit = min(len(df_raw), max_scan)
    best_row, best_score = 0, -float("inf")

    for i in range(scan_limit):
        row = df_raw.iloc[i]
        non_null = [v for v in row if pd.notna(v) and str(v).strip() != ""]
        n_non_null = len(non_null)

        # Skip sparse rows (title / banner / blank)
        if n_non_null < max(2, n_cols * 0.25):
            continue

        # Text ratio — column names are almost always strings
        n_text = sum(1 for v in non_null if isinstance(v, str))
        text_ratio = n_text / n_non_null

        # Uniqueness — duplicate header names are uncommon
        unique_count = len({str(v) for v in non_null})
        uniqueness = unique_count / n_non_null

        # Density — fraction of all columns that have values
        density = n_non_null / n_cols

        score = text_ratio * 3.0 + uniqueness * 2.0 + density * 1.0

        # Bonus when the next row has numeric values (data typically follows header)
        if i + 1 < len(df_raw):
            next_non_null = [v for v in df_raw.iloc[i + 1] if pd.notna(v)]
            next_numeric = sum(1 for v in next_non_null if isinstance(v, (int, float, complex)))
            if next_non_null and next_numeric / len(next_non_null) > 0.1:
                score += 2.0

        if score > best_score:
            best_score = score
            best_row = i

    return best_row


def _read_sheet_with_header_detection(
    source: Any,
    sheet_name: str,
    raw_ws: Any = None,
    engine: str | None = None,
    max_scan: int = 30,
) -> pd.DataFrame:
    """Read a single worksheet with automatic header-row detection.

    Reads the sheet once with ``header=None``, detects the true header row,
    then rebuilds the DataFrame with correct column names and no leading junk
    rows.  The *source* argument is either a :class:`pathlib.Path` or a
    seekable :class:`io.BytesIO` buffer (which is ``seek(0)``-ed before each
    read call).
    """
    read_kwargs: dict[str, Any] = {"sheet_name": sheet_name, "header": None}
    if engine:
        read_kwargs["engine"] = engine

    # First read: no header assumption
    if hasattr(source, "seek"):
        source.seek(0)
    df_raw = pd.read_excel(source, **read_kwargs)

    if df_raw.empty:
        return df_raw

    h = detect_header_row(raw_ws, df_raw, max_scan=max_scan)

    if h == 0:
        # Row 0 is the header — pandas default; re-read properly to get dtypes
        read_kwargs["header"] = 0
        if hasattr(source, "seek"):
            source.seek(0)
        return pd.read_excel(source, **read_kwargs)

    logger.debug(
        "Sheet '%s': header row detected at row %d (1-based: %d)",
        sheet_name, h, h + 1,
    )
    # Rebuild: use detected row as column names, drop all rows up to and
    # including it, reset index.
    col_names = [
        str(v) if pd.notna(v) and str(v).strip() else f"Unnamed:{i}"
        for i, v in enumerate(df_raw.iloc[h])
    ]
    df = df_raw.iloc[h + 1:].copy()
    df.columns = col_names
    return df.reset_index(drop=True)


def load_workbook(path: str | Path) -> WorkbookBundle:
    """Parse every sheet of *path* into a :class:`WorkbookBundle`.

    Supported formats: ``.xlsx``, ``.xlsm``, ``.xls``, ``.csv``.

    Args:
        path: Absolute or relative path to the file.

    Returns:
        A :class:`WorkbookBundle` with one DataFrame per sheet.

    Raises:
        FileNotFoundError: If *path* does not exist.
        ValueError: If the file extension is not supported.
    """
    path = Path(path)

    if not path.exists():
        raise FileNotFoundError(f"File not found: {path}")

    ext = path.suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise ValueError(
            f"Unsupported file type '{ext}'. "
            f"Supported: {', '.join(sorted(SUPPORTED_EXTENSIONS))}"
        )

    logger.info("Loading workbook: %s", path)

    try:
        if ext == ".csv":
            return _load_csv(path)
        elif ext == ".xls":
            return _load_xls(path)
        else:
            return _load_xlsx(path)
    except Exception as exc:
        logger.error("Failed to load %s: %s", path, exc)
        raise


def load_workbook_from_bytes(data: bytes, file_name: str) -> WorkbookBundle:
    """Load an Excel workbook from raw bytes (e.g. a Streamlit upload).

    Args:
        data: Raw bytes of the .xlsx / .xlsm file.
        file_name: Original file name (used as the bundle identifier).

    Returns:
        A :class:`WorkbookBundle` with one DataFrame per sheet.
    """
    import io
    import openpyxl

    buf = io.BytesIO(data)
    raw_wb = openpyxl.load_workbook(buf, data_only=False)
    sheets: dict[str, pd.DataFrame] = {}
    for sheet_name in raw_wb.sheetnames:
        try:
            raw_ws = raw_wb[sheet_name]
        except Exception:
            raw_ws = None
        df = _read_sheet_with_header_detection(buf, sheet_name, raw_ws=raw_ws, engine="openpyxl")
        sheets[sheet_name] = df
        logger.debug("  Sheet '%s': %d rows × %d cols", sheet_name, len(df), len(df.columns))

    fake_path = Path(file_name)
    logger.info("Loaded %d sheet(s) from bytes (%s)", len(sheets), file_name)
    return WorkbookBundle(source_path=fake_path, sheets=sheets, _raw_wb=raw_wb)


def load_many(paths: list[str | Path]) -> list[WorkbookBundle]:
    """Load multiple files, skipping those that fail with a logged warning.

    Args:
        paths: Iterable of file paths.

    Returns:
        List of successfully loaded :class:`WorkbookBundle` objects.
    """
    bundles: list[WorkbookBundle] = []
    for p in paths:
        try:
            bundles.append(load_workbook(p))
        except Exception as exc:
            logger.warning("Skipping %s — %s", p, exc)
    return bundles


# ---------------------------------------------------------------------------
# Private helpers
# ---------------------------------------------------------------------------

def _load_xlsx(path: Path) -> WorkbookBundle:
    """Load .xlsx / .xlsm using openpyxl (preserves formula metadata)."""
    import openpyxl  # local import keeps module importable without openpyxl installed

    raw_wb = openpyxl.load_workbook(path, data_only=False)
    sheets: dict[str, pd.DataFrame] = {}

    for sheet_name in raw_wb.sheetnames:
        try:
            raw_ws = raw_wb[sheet_name]
        except Exception:
            raw_ws = None
        df = _read_sheet_with_header_detection(path, sheet_name, raw_ws=raw_ws, engine="openpyxl")
        sheets[sheet_name] = df
        logger.debug("  Sheet '%s': %d rows × %d cols", sheet_name, len(df), len(df.columns))

    logger.info("Loaded %d sheet(s) from %s", len(sheets), path.name)
    return WorkbookBundle(source_path=path, sheets=sheets, _raw_wb=raw_wb)


def _load_xls(path: Path) -> WorkbookBundle:
    """Load legacy .xls using xlrd."""
    xl = pd.ExcelFile(path, engine="xlrd")
    sheets: dict[str, pd.DataFrame] = {}
    for name in xl.sheet_names:
        df = _read_sheet_with_header_detection(path, name, raw_ws=None, engine="xlrd")
        sheets[name] = df
    logger.info("Loaded %d sheet(s) from %s (xls)", len(sheets), path.name)
    return WorkbookBundle(source_path=path, sheets=sheets)


def _load_csv(path: Path) -> WorkbookBundle:
    """Load a CSV as a single-sheet bundle (sheet name = file stem)."""
    df = pd.read_csv(path)
    sheet_name = path.stem
    logger.info("Loaded CSV '%s': %d rows × %d cols", path.name, len(df), len(df.columns))
    return WorkbookBundle(source_path=path, sheets={sheet_name: df})
