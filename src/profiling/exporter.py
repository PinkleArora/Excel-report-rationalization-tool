"""Export profiling results to a multi-sheet Excel workbook."""

from __future__ import annotations

import logging
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook as openpyxl_load
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from src.profiling.metadata import WorkbookMetadata
from src.profiling.profiler import build_summary_dataframes

logger = logging.getLogger(__name__)

# House-style colours
_HEADER_FILL = "1F4E79"   # dark blue
_HEADER_FONT = "FFFFFF"   # white
_ALT_ROW_FILL = "D6E4F0"  # light blue

OUTPUT_SHEETS = {
    "Workbook_Summary": 0,
    "Sheet_Summary": 1,
    "Column_Summary": 2,
}


def export_profiling_results(
    workbook_metas: list[WorkbookMetadata],
    output_path: str | Path | None = None,
) -> Path:
    """Write a formatted 3-sheet profiling workbook.

    Args:
        workbook_metas: List of :class:`WorkbookMetadata` objects produced by
            :func:`~src.profiling.profiler.profile_workbook`.
        output_path: Destination ``.xlsx`` path.  Defaults to
            ``output/profiling_summary.xlsx`` relative to the working directory.

    Returns:
        The resolved path of the written file.

    Raises:
        ValueError: If *workbook_metas* is empty.
        OSError: If the output directory cannot be created.
    """
    if not workbook_metas:
        raise ValueError("No workbook metadata provided — nothing to export.")

    output_path = Path(output_path) if output_path else Path("output") / "profiling_summary.xlsx"
    output_path.parent.mkdir(parents=True, exist_ok=True)

    wb_df, sheet_df, col_df = build_summary_dataframes(workbook_metas)

    logger.info("Writing profiling summary to %s", output_path)

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        wb_df.to_excel(writer, sheet_name="Workbook_Summary", index=False)
        sheet_df.to_excel(writer, sheet_name="Sheet_Summary", index=False)
        col_df.to_excel(writer, sheet_name="Column_Summary", index=False)

    _apply_styles(output_path)
    logger.info("Export complete: %s", output_path.resolve())
    return output_path.resolve()


# ---------------------------------------------------------------------------
# Styling
# ---------------------------------------------------------------------------

def _apply_styles(path: Path) -> None:
    wb = openpyxl_load(path)
    for ws in wb.worksheets:
        _style_header(ws)
        _auto_width(ws)
        _freeze_and_filter(ws)
        _alternate_rows(ws)
    wb.save(path)


def _style_header(ws) -> None:
    fill = PatternFill(fill_type="solid", fgColor=_HEADER_FILL)
    font = Font(bold=True, color=_HEADER_FONT, name="Calibri", size=11)
    for cell in ws[1]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=False)


def _auto_width(ws) -> None:
    for col_cells in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col_cells[0].column)
        for cell in col_cells:
            try:
                cell_len = len(str(cell.value)) if cell.value is not None else 0
                max_len = max(max_len, cell_len)
            except Exception:
                pass
        ws.column_dimensions[col_letter].width = min(max_len + 4, 60)


def _freeze_and_filter(ws) -> None:
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def _alternate_rows(ws) -> None:
    alt_fill = PatternFill(fill_type="solid", fgColor=_ALT_ROW_FILL)
    for i, row in enumerate(ws.iter_rows(min_row=2), start=2):
        if i % 2 == 0:
            for cell in row:
                if cell.fill.fill_type == "none" or cell.fill.fgColor.rgb in ("00000000", "FFFFFFFF"):
                    cell.fill = alt_fill
