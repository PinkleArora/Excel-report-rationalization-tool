"""Visual rationalization diagram builder.

Produces two diagrams as PNG and SVG bytes, and can embed them into an
openpyxl Workbook as new worksheets.

  build_rationalization_flow()  — 4-stage left→right flowchart
  build_dependency_diagram()    — source column → master column → KPI map
  embed_diagrams_in_workbook()  — adds both as embedded images in Excel sheets
"""

from __future__ import annotations

import io
import logging
import textwrap
from collections import defaultdict

import matplotlib
matplotlib.use("Agg")  # non-interactive backend; must come before pyplot import
import matplotlib.patches as mpatches
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch
from matplotlib.path import Path
import matplotlib.patheffects as pe
import numpy as np

from src.guided_rationalization.config import RationalizationConfig
from src.guided_rationalization.kpi_analyzer import KpiAnalysisResult
from src.guided_rationalization.source_analyzer import SourceAnalysisResult

logger = logging.getLogger(__name__)

# ── Colour palette ────────────────────────────────────────────────────────────
_C_INPUT    = "#1F4E79"   # dark blue — input workbooks
_C_COMMON   = "#2E75B6"   # mid blue  — common/merged columns
_C_UNIQUE   = "#5B9BD5"   # light blue — unique columns
_C_EXCLUDED = "#A9A9A9"   # grey      — excluded columns
_C_MASTER   = "#375623"   # dark green — master source data
_C_KPI      = "#7030A0"   # purple    — KPI tabs
_C_FUTURE   = "#833C00"   # brown     — future-state workbook
_C_ARROW    = "#595959"   # dark grey — arrows
_C_TEXT_LT  = "#FFFFFF"   # white text on dark backgrounds
_C_TEXT_DK  = "#1A1A1A"   # dark text on light backgrounds
_C_BG       = "#F8F9FA"   # page background

_MAX_COL_ITEMS  = 12   # max column names shown per box before "… N more"
_MAX_NAME_LEN   = 32   # max characters per name before truncation


def _trunc(s: str, n: int = _MAX_NAME_LEN) -> str:
    return s if len(s) <= n else s[: n - 1] + "…"


def _col_bullet(name: str, n: int = _MAX_NAME_LEN) -> str:
    return "• " + _trunc(name, n)


# ── Rationalization Flow ──────────────────────────────────────────────────────

def build_rationalization_flow(
    source_result: SourceAnalysisResult,
    kpi_result: KpiAnalysisResult,
    config: RationalizationConfig,
    master_df=None,
) -> tuple[bytes, bytes]:
    """4-stage left→right rationalization flow diagram.

    Returns:
        (png_bytes, svg_bytes)
    """
    # ── Gather data ──────────────────────────────────────────────────────────
    wb_cfgs = config.workbook_configs

    # Stage 1 — Input workbooks
    stage1_boxes = []
    for wb_cfg in wb_cfgs:
        lines = [
            _trunc(wb_cfg.workbook_name, 30),
            f"Source: {_trunc(wb_cfg.source_tab, 24)}",
        ]
        for kt in wb_cfg.kpi_tabs:
            lines.append(f"KPI: {_trunc(kt, 26)}")
        stage1_boxes.append(lines)

    # Stage 2 — Source consolidation
    common_names  = [p.canonical_name for p in source_result.common_columns]
    unique_names  = [p.canonical_name for p in source_result.unique_columns]
    similar_names = [p.canonical_name for p in source_result.similar_columns]
    excl_set      = set(source_result.excluded_columns)
    excl_count    = len(excl_set)

    def _fmt_list(items: list[str], label: str) -> list[str]:
        out = [f"── {label} ({len(items)}) ──"]
        for name in items[:_MAX_COL_ITEMS]:
            out.append(_col_bullet(name))
        if len(items) > _MAX_COL_ITEMS:
            out.append(f"  … {len(items) - _MAX_COL_ITEMS} more")
        return out

    stage2_lines: list[str] = []
    if common_names:
        stage2_lines += _fmt_list(common_names, "Common (merged)")
    if unique_names:
        if stage2_lines:
            stage2_lines.append("")
        stage2_lines += _fmt_list(unique_names, "Unique (retained)")
    if similar_names:
        if stage2_lines:
            stage2_lines.append("")
        stage2_lines += _fmt_list(similar_names, "Similar (kept separate)")
    if excl_count:
        if stage2_lines:
            stage2_lines.append("")
        stage2_lines.append(f"── Excluded ({excl_count}) ──")
        stage2_lines.append("Not referenced by any KPI")

    # Stage 3 — Master Source Data
    retained = [
        c for c in {canonical for canonical in source_result.column_mapping.values()}
        if any(
            (wb, tab, orig) not in excl_set
            for (wb, tab, orig), can in source_result.column_mapping.items()
            if can == c
        )
    ]
    master_rows = len(master_df) if master_df is not None and not master_df.empty else "?"
    master_cols = len(retained)
    stage3_lines = [
        f"{master_rows} rows  ×  {master_cols} columns",
        "",
        "── Retained Columns ──",
    ]
    for name in sorted(retained)[:_MAX_COL_ITEMS]:
        stage3_lines.append(_col_bullet(name))
    if len(retained) > _MAX_COL_ITEMS:
        stage3_lines.append(f"  … {len(retained) - _MAX_COL_ITEMS} more")

    # Stage 4 — Future-state workbook
    stage4_lines = [config.future_source_tab_name]
    for wb_cfg in wb_cfgs:
        for kt in wb_cfg.kpi_tabs:
            tab_name = config.kpi_tab_name_for(wb_cfg.workbook_name, kt)
            stage4_lines.append(_trunc(tab_name, 28))

    # ── Layout ───────────────────────────────────────────────────────────────
    n_s1    = max(len(stage1_boxes), 1)
    n_s2    = max(len(stage2_lines), 1)
    n_s3    = max(len(stage3_lines), 1)
    n_s4    = max(len(stage4_lines), 1)

    row_h   = 0.28   # height per text row (inches)
    pad     = 0.35   # vertical padding inside box
    hdr_h   = 0.55   # header band height

    # box heights
    s1_h  = sum(hdr_h + pad + len(b) * row_h + pad for b in stage1_boxes) + 0.2
    s2_h  = hdr_h + pad + n_s2 * row_h + pad
    s3_h  = hdr_h + pad + n_s3 * row_h + pad
    s4_h  = hdr_h + pad + n_s4 * row_h + pad

    fig_h   = max(s1_h, s2_h, s3_h, s4_h) + 2.0
    fig_w   = 22.0

    fig, ax = plt.subplots(figsize=(fig_w, fig_h))
    ax.set_xlim(0, fig_w)
    ax.set_ylim(0, fig_h)
    ax.axis("off")
    fig.patch.set_facecolor(_C_BG)

    # Title
    ax.text(
        fig_w / 2, fig_h - 0.45,
        "Rationalization Flow",
        ha="center", va="center",
        fontsize=17, fontweight="bold", color=_C_INPUT,
    )

    # Stage x-centres
    col_w   = fig_w / 4
    cx      = [col_w * 0.5, col_w * 1.5, col_w * 2.5, col_w * 3.5]
    box_w   = col_w - 0.6
    cy      = fig_h / 2 - 0.4   # vertical centre for single boxes
    arrow_y = fig_h / 2 - 0.4

    stage_labels = [
        ("1  Input Workbooks",     _C_INPUT,   cx[0]),
        ("2  Source Consolidation", _C_COMMON,  cx[1]),
        ("3  Master Source Data",   _C_MASTER,  cx[2]),
        ("4  Future-State Workbook", _C_FUTURE, cx[3]),
    ]

    def _draw_box(ax_, x, y_bottom, width, height, header, lines,
                  hdr_color, body_bg="#FFFFFF", hdr_font_size=10):
        # Body
        body = FancyBboxPatch(
            (x - width / 2, y_bottom), width, height,
            boxstyle="round,pad=0.02", linewidth=1.2,
            edgecolor=hdr_color, facecolor=body_bg,
        )
        ax_.add_patch(body)
        # Header band
        hdr = FancyBboxPatch(
            (x - width / 2, y_bottom + height - hdr_h), width, hdr_h,
            boxstyle="round,pad=0.02", linewidth=0,
            facecolor=hdr_color, clip_on=True,
        )
        ax_.add_patch(hdr)
        ax_.text(
            x, y_bottom + height - hdr_h / 2, header,
            ha="center", va="center",
            fontsize=hdr_font_size, fontweight="bold", color=_C_TEXT_LT,
            clip_on=True,
        )
        # Content lines
        y_text = y_bottom + height - hdr_h - pad
        for line in lines:
            if line.startswith("──"):
                ax_.text(
                    x - width / 2 + 0.15, y_text, line,
                    ha="left", va="top", fontsize=7.5, fontweight="bold",
                    color=hdr_color, clip_on=True,
                )
            else:
                ax_.text(
                    x - width / 2 + 0.15, y_text, line,
                    ha="left", va="top", fontsize=7.5, color=_C_TEXT_DK,
                    clip_on=True,
                )
            y_text -= row_h
        return (y_bottom, y_bottom + height)   # bottom, top y coords

    # Stage 1 — stacked boxes, one per workbook
    s1_total_h = sum(
        hdr_h + pad + len(b) * row_h + pad for b in stage1_boxes
    )
    s1_gap = 0.15
    s1_start = cy - s1_total_h / 2 - (len(stage1_boxes) - 1) * s1_gap / 2
    s1_ymids = []
    y_cur = s1_start
    for box_lines in stage1_boxes:
        bh = hdr_h + pad + len(box_lines) * row_h + pad
        header = box_lines[0]
        content = box_lines[1:]
        bot, top = _draw_box(ax, cx[0], y_cur, box_w, bh,
                             header, content, _C_INPUT, body_bg="#EBF3FA")
        s1_ymids.append((bot + top) / 2)
        y_cur += bh + s1_gap

    arrow_y_s1 = (s1_start + y_cur - s1_gap) / 2

    # Stage 2 — single box
    s2_box_h = s2_h
    s2_y_bot = cy - s2_box_h / 2
    bot2, top2 = _draw_box(ax, cx[1], s2_y_bot, box_w, s2_box_h,
                            "Source Consolidation", stage2_lines,
                            _C_COMMON, body_bg="#EBF3FA")

    # Stage 3 — single box
    s3_box_h = s3_h
    s3_y_bot = cy - s3_box_h / 2
    bot3, top3 = _draw_box(ax, cx[2], s3_y_bot, box_w, s3_box_h,
                            "Master Source Data", stage3_lines,
                            _C_MASTER, body_bg="#EBF2E9")

    # Stage 4 — single box showing tab hierarchy
    s4_box_h = s4_h
    s4_y_bot = cy - s4_box_h / 2
    bot4, top4 = _draw_box(ax, cx[3], s4_y_bot, box_w, s4_box_h,
                            "Future-State Workbook", stage4_lines,
                            _C_FUTURE, body_bg="#FDF3EC")

    # ── Arrows ────────────────────────────────────────────────────────────────
    arrow_props = dict(
        arrowstyle="-|>", color=_C_ARROW,
        lw=1.8, mutation_scale=18,
        connectionstyle="arc3,rad=0",
    )

    def _arrow(ax_, x0, y0, x1, y1, **kw):
        props = {**arrow_props, **kw}
        ax_.annotate(
            "", xy=(x1, y1), xytext=(x0, y0),
            arrowprops=dict(
                arrowstyle="-|>", color=props.get("color", _C_ARROW),
                lw=props.get("lw", 1.8), mutation_scale=18,
                connectionstyle=props.get("connectionstyle", "arc3,rad=0"),
            ),
        )

    mid12_y = (bot2 + top2) / 2
    mid23_y = (bot3 + top3) / 2
    mid34_y = (bot4 + top4) / 2

    # S1 boxes → S2
    for ymid in s1_ymids:
        _arrow(ax, cx[0] + box_w / 2, ymid, cx[1] - box_w / 2, mid12_y)

    # S2 → S3
    _arrow(ax, cx[1] + box_w / 2, mid12_y, cx[2] - box_w / 2, mid23_y)

    # S3 → S4
    _arrow(ax, cx[2] + box_w / 2, mid23_y, cx[3] - box_w / 2, mid34_y)

    # Stage labels at top of each column
    for label, colour, x in stage_labels:
        ax.text(
            x, fig_h - 0.9, label,
            ha="center", va="center",
            fontsize=9.5, fontweight="bold", color=colour,
        )

    plt.tight_layout(pad=0.3)

    png_buf = io.BytesIO()
    fig.savefig(png_buf, format="png", dpi=150, bbox_inches="tight",
                facecolor=_C_BG)
    png_bytes = png_buf.getvalue()

    svg_buf = io.BytesIO()
    fig.savefig(svg_buf, format="svg", bbox_inches="tight", facecolor=_C_BG)
    svg_bytes = svg_buf.getvalue()

    plt.close(fig)
    return png_bytes, svg_bytes


# ── Dependency Diagram ────────────────────────────────────────────────────────

def build_dependency_diagram(
    source_result: SourceAnalysisResult,
    kpi_result: KpiAnalysisResult,
    config: RationalizationConfig,
) -> tuple[bytes, bytes]:
    """Source column → Master column → KPI tab dependency diagram.

    Only KPI-referenced columns are shown in the master column lane to
    keep the diagram readable.  Source entries are grouped by workbook.

    Returns:
        (png_bytes, svg_bytes)
    """
    # ── Data preparation ──────────────────────────────────────────────────────
    # Canonical → KPI tabs (future-state tab names)
    canonical_to_kpi_tabs: dict[str, list[str]] = defaultdict(list)
    for dep in kpi_result.dependencies:
        future_tab = config.kpi_tab_name_for(dep.workbook_name, dep.kpi_tab)
        for canonical in dep.canonical_source_columns:
            if future_tab not in canonical_to_kpi_tabs[canonical]:
                canonical_to_kpi_tabs[canonical].append(future_tab)

    # Only keep canonicals that have KPI connections
    ref_canonicals = sorted(kpi_result.referenced_canonicals)

    # Collect unique KPI tabs in order
    all_kpi_tabs: list[str] = []
    for dep in kpi_result.dependencies:
        tab = config.kpi_tab_name_for(dep.workbook_name, dep.kpi_tab)
        if tab not in all_kpi_tabs:
            all_kpi_tabs.append(tab)

    # Source entries: (workbook, source_tab, original_col) for ref'd canonicals
    # Group by workbook — use canonical for display in centre lane
    # Left lane: workbook → [original col names that map to ref'd canonicals]
    excl_set = set(source_result.excluded_columns)
    wb_to_src_cols: dict[str, list[tuple[str, str]]] = defaultdict(list)
    # (original_col, canonical)
    for (wb, tab, orig), canonical in source_result.column_mapping.items():
        if canonical in kpi_result.referenced_canonicals and (wb, tab, orig) not in excl_set:
            if (orig, canonical) not in wb_to_src_cols[wb]:
                wb_to_src_cols[wb].append((orig, canonical))

    # Sort within each workbook
    for wb in wb_to_src_cols:
        wb_to_src_cols[wb].sort(key=lambda t: t[1])

    workbooks = list(wb_to_src_cols.keys())

    # If nothing to show
    if not ref_canonicals and not workbooks:
        fig, ax = plt.subplots(figsize=(10, 4))
        ax.axis("off")
        ax.text(0.5, 0.5, "No KPI-referenced columns found.",
                ha="center", va="center", fontsize=14)
        png_buf = io.BytesIO(); fig.savefig(png_buf, format="png", dpi=150)
        svg_buf = io.BytesIO(); fig.savefig(svg_buf, format="svg")
        plt.close(fig)
        return png_buf.getvalue(), svg_buf.getvalue()

    # ── Layout constants ──────────────────────────────────────────────────────
    row_h       = 0.38   # height per item row
    box_pad_v   = 0.25   # top/bottom padding inside a box
    hdr_h       = 0.52
    col_gap     = 1.8    # horizontal gap between column groups
    box_w       = 4.8    # width of each item box
    item_h      = 0.34   # height of individual row items

    # Column x-positions
    x_src   = 1.5
    x_mid   = x_src + box_w + col_gap
    x_kpi   = x_mid + box_w + col_gap
    fig_w   = x_kpi + box_w + 1.5

    # Vertical layout — derive total rows per column
    # Left: all source entries (across workbooks, with workbook headers)
    left_items: list[tuple[str, str, str, str]] = []
    # (display_name, canonical, workbook, "header"|"item")
    for wb in workbooks:
        left_items.append((wb, "", wb, "header"))
        for orig, canonical in wb_to_src_cols[wb]:
            left_items.append((orig, canonical, wb, "item"))

    n_left   = len(left_items)
    n_middle = len(ref_canonicals)
    n_right  = len(all_kpi_tabs)

    total_rows = max(n_left, n_middle, n_right, 1)
    fig_h = total_rows * row_h + 3.0
    fig_h = max(fig_h, 8.0)

    fig, ax = plt.subplots(figsize=(fig_w, fig_h))
    ax.set_xlim(0, fig_w)
    ax.set_ylim(0, fig_h)
    ax.axis("off")
    fig.patch.set_facecolor(_C_BG)

    # Title
    ax.text(
        fig_w / 2, fig_h - 0.45,
        "Column Dependency Diagram  (KPI-Referenced Columns Only)",
        ha="center", va="center",
        fontsize=15, fontweight="bold", color=_C_INPUT,
    )

    # Column headers
    for label, colour, x in [
        ("Source Columns", _C_INPUT,  x_src  + box_w / 2),
        ("Master Source Columns", _C_MASTER, x_mid  + box_w / 2),
        ("KPI Tabs", _C_KPI,   x_kpi + box_w / 2),
    ]:
        ax.text(
            x, fig_h - 0.95, label,
            ha="center", va="center",
            fontsize=11, fontweight="bold", color=colour,
        )
        ax.plot(
            [x - box_w / 2 + 0.1, x + box_w / 2 - 0.1],
            [fig_h - 1.15, fig_h - 1.15],
            color=colour, lw=2,
        )

    content_top = fig_h - 1.45

    # ── Draw items and record y-centres for connection lines ─────────────────
    # Workbook colours (cycle through palette)
    wb_colours = ["#1F4E79", "#2E75B6", "#5B9BD5", "#375623", "#7030A0", "#833C00"]
    wb_colour_map = {wb: wb_colours[i % len(wb_colours)] for i, wb in enumerate(workbooks)}

    def _item_box(ax_, x, y_centre, width, text, bg_color, text_color=_C_TEXT_LT,
                  is_header=False, fontsize=8.5):
        h = item_h * (1.3 if is_header else 1.0)
        patch = FancyBboxPatch(
            (x, y_centre - h / 2), width, h,
            boxstyle="round,pad=0.04",
            linewidth=0.8 if not is_header else 1.5,
            edgecolor=bg_color,
            facecolor=bg_color if is_header else bg_color + "28",   # semi-transparent for items
        )
        ax_.add_patch(patch)
        tcolor = text_color if is_header else _C_TEXT_DK
        ax_.text(
            x + 0.12, y_centre, _trunc(text, 36 if width > 4 else 28),
            ha="left", va="center", fontsize=fontsize,
            color=tcolor, fontweight="bold" if is_header else "normal",
        )
        return y_centre

    # LEFT column y-positions
    left_y: dict[tuple[str, str], float] = {}   # (orig, canonical) → y
    y_cur = content_top
    for display, canonical, wb, kind in left_items:
        colour = wb_colour_map.get(wb, _C_INPUT)
        if kind == "header":
            _item_box(ax, x_src, y_cur, box_w, f"[WB]  {_trunc(display, 30)}", colour,
                      is_header=True, fontsize=9)
        else:
            _item_box(ax, x_src, y_cur, box_w, display, colour,
                      is_header=False, fontsize=8.5)
            left_y[(display, canonical)] = y_cur
        y_cur -= row_h

    # MIDDLE column y-positions — evenly spaced across same content height
    mid_total_h = content_top - (content_top - n_left * row_h)
    if n_middle > 0:
        mid_spacing = (n_left * row_h) / max(n_middle, 1)
    else:
        mid_spacing = row_h
    mid_y: dict[str, float] = {}
    y_cur_mid = content_top - mid_spacing / 2
    for canonical in ref_canonicals:
        # Colour by match_class
        profile = next(
            (p for p in source_result.column_profiles if p.canonical_name == canonical), None
        )
        if profile:
            mc = profile.match_class
            colour = _C_COMMON if mc == "common" else (_C_UNIQUE if mc == "unique" else _C_KPI)
        else:
            colour = _C_MASTER
        _item_box(ax, x_mid, y_cur_mid, box_w, canonical, colour,
                  is_header=False, fontsize=8)
        mid_y[canonical] = y_cur_mid
        y_cur_mid -= mid_spacing

    # RIGHT column y-positions
    if n_right > 0:
        kpi_spacing = (n_left * row_h) / max(n_right, 1)
    else:
        kpi_spacing = row_h
    kpi_y: dict[str, float] = {}
    y_cur_kpi = content_top - kpi_spacing / 2
    for tab in all_kpi_tabs:
        _item_box(ax, x_kpi, y_cur_kpi, box_w, tab, _C_KPI,
                  is_header=True, fontsize=9)
        kpi_y[tab] = y_cur_kpi
        y_cur_kpi -= kpi_spacing

    # ── Connection lines ──────────────────────────────────────────────────────
    def _bezier(ax_, x0, y0, x1, y1, colour, alpha=0.35, lw=1.0):
        """Draw a smooth bezier curve from (x0,y0) to (x1,y1)."""
        mx = (x0 + x1) / 2
        verts = [(x0, y0), (mx, y0), (mx, y1), (x1, y1)]
        codes = [Path.MOVETO, Path.CURVE4, Path.CURVE4, Path.CURVE4]
        path = Path(verts, codes)
        patch = mpatches.PathPatch(
            path, facecolor="none", edgecolor=colour,
            lw=lw, alpha=alpha,
        )
        ax_.add_patch(patch)

    # Left → Middle (source → canonical)
    for (orig, canonical), y_left in left_y.items():
        if canonical in mid_y:
            wb_entries = [wb for wb, entries in wb_to_src_cols.items()
                          if any(o == orig and c == canonical for o, c in entries)]
            colour = wb_colour_map.get(wb_entries[0], _C_INPUT) if wb_entries else _C_INPUT
            _bezier(ax, x_src + box_w, y_left, x_mid, mid_y[canonical],
                    colour, alpha=0.4, lw=1.2)

    # Middle → Right (canonical → KPI tabs)
    for canonical in ref_canonicals:
        if canonical not in mid_y:
            continue
        for kpi_tab in canonical_to_kpi_tabs.get(canonical, []):
            if kpi_tab in kpi_y:
                _bezier(ax, x_mid + box_w, mid_y[canonical],
                        x_kpi, kpi_y[kpi_tab], _C_KPI, alpha=0.3, lw=1.1)

    # Legend
    legend_y = 0.55
    legend_items = [
        ("Common column",       _C_COMMON),
        ("Unique column",       _C_UNIQUE),
        ("Similar column",      _C_KPI),
        ("Workbook (header)",   wb_colours[0]),
    ]
    ax.text(0.3, legend_y + 0.1, "Legend:", fontsize=8, fontweight="bold", color=_C_TEXT_DK)
    for i, (label, colour) in enumerate(legend_items):
        x_leg = 1.2 + i * 4.5
        patch = FancyBboxPatch(
            (x_leg, legend_y - 0.13), 0.3, 0.25,
            boxstyle="round,pad=0.03", facecolor=colour, edgecolor=colour,
        )
        ax.add_patch(patch)
        ax.text(x_leg + 0.38, legend_y, label,
                fontsize=8, va="center", color=_C_TEXT_DK)

    plt.tight_layout(pad=0.3)

    png_buf = io.BytesIO()
    fig.savefig(png_buf, format="png", dpi=150, bbox_inches="tight", facecolor=_C_BG)
    png_bytes = png_buf.getvalue()

    svg_buf = io.BytesIO()
    fig.savefig(svg_buf, format="svg", bbox_inches="tight", facecolor=_C_BG)
    svg_bytes = svg_buf.getvalue()

    plt.close(fig)
    return png_bytes, svg_bytes


# ── Excel embedding ───────────────────────────────────────────────────────────

def embed_diagrams_in_workbook(
    wb,
    flow_png: bytes,
    dep_png: bytes,
    flow_sheet_name: str = "Rationalization_Flow",
    dep_sheet_name: str = "Dependency_Diagram",
) -> None:
    """Embed PNG diagrams as images in two new worksheets of an openpyxl Workbook.

    The sheets are inserted at position 0 so they appear first.
    """
    from openpyxl.drawing.image import Image as XLImage

    for sheet_name, png_bytes, title in [
        (flow_sheet_name, flow_png, "Rationalization Flow Diagram"),
        (dep_sheet_name,  dep_png,  "Column Dependency Diagram"),
    ]:
        ws = wb.create_sheet(sheet_name, 0)
        ws.sheet_view.showGridLines = False

        # Title cell
        from openpyxl.styles import Font, Alignment, PatternFill
        title_cell = ws.cell(row=1, column=1, value=title)
        title_cell.font = Font(bold=True, size=14, color="1F4E79")
        title_cell.alignment = Alignment(horizontal="left")
        ws.row_dimensions[1].height = 22

        # Embed image starting at B3
        img_stream = io.BytesIO(png_bytes)
        xl_img = XLImage(img_stream)
        # Scale to fit reasonably in the sheet (max ~1400px wide)
        if xl_img.width > 1400:
            scale = 1400 / xl_img.width
            xl_img.width  = int(xl_img.width  * scale)
            xl_img.height = int(xl_img.height * scale)
        ws.add_image(xl_img, "B3")

        # Widen columns so the image has breathing room
        for col_letter in "ABCDEFGHIJKLMNOPQRSTUVWXYZ":
            ws.column_dimensions[col_letter].width = 12


# ── Convenience: build both and return bytes ──────────────────────────────────

def build_all_diagrams(
    source_result: SourceAnalysisResult,
    kpi_result: KpiAnalysisResult,
    config: RationalizationConfig,
    master_df=None,
) -> dict[str, bytes]:
    """Build both diagrams and return a dict of all byte outputs.

    Keys: 'flow_png', 'flow_svg', 'dep_png', 'dep_svg'
    """
    try:
        flow_png, flow_svg = build_rationalization_flow(
            source_result, kpi_result, config, master_df
        )
    except Exception as exc:
        logger.warning("build_rationalization_flow failed: %s", exc)
        flow_png = flow_svg = b""

    try:
        dep_png, dep_svg = build_dependency_diagram(source_result, kpi_result, config)
    except Exception as exc:
        logger.warning("build_dependency_diagram failed: %s", exc)
        dep_png = dep_svg = b""

    return {
        "flow_png": flow_png,
        "flow_svg": flow_svg,
        "dep_png":  dep_png,
        "dep_svg":  dep_svg,
    }
