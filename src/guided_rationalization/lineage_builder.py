"""Data Lineage builders.

Produces three DataFrames and one assembled workbook from the outputs of
source_analyzer and kpi_analyzer:

  Data_Lineage      — KPI formula → source column traceability
  Column_Lineage    — Master column → original column provenance
  KPI_Dependencies  — Complete KPI-to-source dependency map (summary view)
"""

from __future__ import annotations

import io
import logging
from collections import defaultdict

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from src.guided_rationalization.config import RationalizationConfig
from src.guided_rationalization.kpi_analyzer import KpiAnalysisResult
from src.guided_rationalization.source_analyzer import SourceAnalysisResult
from src.schema_matching.normalizer import normalize_column_name

logger = logging.getLogger(__name__)

# ── Styling ──────────────────────────────────────────────────────────────────
_HEADER_FILL  = PatternFill("solid", fgColor="1F4E79")
_HEADER_FONT  = Font(color="FFFFFF", bold=True)
_ALT_FILL     = PatternFill("solid", fgColor="D6E4F0")
_HIGH_FILL    = PatternFill("solid", fgColor="C6EFCE")   # green
_MED_FILL     = PatternFill("solid", fgColor="FFEB9C")   # amber
_LOW_FILL     = PatternFill("solid", fgColor="FFC7CE")   # red


# ── Internal helpers ─────────────────────────────────────────────────────────

def _build_reverse_mapping(
    source_result: SourceAnalysisResult,
) -> dict[str, list[tuple[str, str, str]]]:
    """Return canonical → [(workbook, tab, original_col), …]."""
    reverse: dict[str, list[tuple[str, str, str]]] = defaultdict(list)
    for (wb, tab, orig), canonical in source_result.column_mapping.items():
        reverse[canonical].append((wb, tab, orig))
    return dict(reverse)


def _transformation_label(
    canonical: str,
    source_entries: list[tuple[str, str, str]],
    match_class: str,
) -> str:
    """Plain-English description of how a canonical column was formed."""
    if match_class == "common":
        wbs = sorted({wb for wb, _, _ in source_entries})
        if len(wbs) > 1:
            return f"Merged from {len(wbs)} workbooks (identical normalized name)"
        return "Column standardized and retained"
    if match_class == "similar":
        return "Kept separate (similar name to other columns — distinct business field)"
    # unique
    orig_names = sorted({orig for _, _, orig in source_entries})
    if len(orig_names) == 1 and normalize_column_name(orig_names[0]) == canonical:
        return "Column standardized and retained"
    return "Retained as unique column"


def _kpi_confidence(dep) -> str:
    """Confidence string for a KpiDependency."""
    if not dep.canonical_source_columns:
        return "40%"
    if not dep.refs_source_tab_only:
        return "75%"
    return "100%"


def _kpi_notes(dep) -> str:
    if not dep.canonical_source_columns:
        return "No source columns resolved — formula may reference non-source tab."
    if not dep.refs_source_tab_only:
        return "Formula references sheet(s) other than the configured source tab."
    return ""


# ── Deliverable 1: Data_Lineage ───────────────────────────────────────────────

def build_data_lineage_df(
    source_result: SourceAnalysisResult,
    kpi_result: KpiAnalysisResult,
    config: RationalizationConfig,
) -> pd.DataFrame:
    """One row per KPI formula — full KPI → source column traceability.

    Columns:
        KPI Workbook, KPI Tab, KPI Cell Address, KPI Label,
        Aggregate Function, Original Formula,
        Source Workbook, Source Tab,
        Original Source Column(s), Master Source Column(s),
        Transformation Applied, Confidence Score, Notes
    """
    reverse = _build_reverse_mapping(source_result)
    rows = []

    for dep in kpi_result.dependencies:
        # Original source columns — reverse-lookup each canonical
        orig_col_parts: list[str] = []
        for canonical in dep.canonical_source_columns:
            entries = reverse.get(canonical, [])
            # Filter to entries from this workbook's source tab
            wb_cfg = config.config_for(dep.workbook_name)
            src_tab = wb_cfg.source_tab if wb_cfg else ""
            wb_entries = [
                orig for wb, tab, orig in entries
                if wb == dep.workbook_name and tab == src_tab
            ]
            if wb_entries:
                orig_col_parts.extend(wb_entries)
            else:
                # fallback: show canonical
                orig_col_parts.append(canonical)

        # Transformation per canonical
        transformations: list[str] = []
        for canonical in dep.canonical_source_columns:
            profile = next(
                (p for p in source_result.column_profiles if p.canonical_name == canonical),
                None,
            )
            if profile:
                transformations.append(
                    _transformation_label(canonical, profile.source_entries, profile.match_class)
                )

        wb_cfg = config.config_for(dep.workbook_name)
        rows.append({
            "KPI Workbook":            dep.workbook_name,
            "KPI Tab":                 dep.kpi_tab,
            "KPI Cell Address":        dep.cell_address,
            "KPI Label":               dep.kpi_label,
            "Aggregate Function":      dep.aggregate_function,
            "Original Formula":        dep.formula,
            "Source Workbook":         dep.workbook_name,
            "Source Tab":              wb_cfg.source_tab if wb_cfg else "",
            "Original Source Column(s)": "; ".join(dict.fromkeys(orig_col_parts)),
            "Master Source Column(s)": "; ".join(dep.canonical_source_columns),
            "Transformation Applied":  "; ".join(dict.fromkeys(transformations)) or "—",
            "Confidence Score":        _kpi_confidence(dep),
            "Notes":                   _kpi_notes(dep),
        })

    if not rows:
        return pd.DataFrame(columns=[
            "KPI Workbook", "KPI Tab", "KPI Cell Address", "KPI Label",
            "Aggregate Function", "Original Formula",
            "Source Workbook", "Source Tab",
            "Original Source Column(s)", "Master Source Column(s)",
            "Transformation Applied", "Confidence Score", "Notes",
        ])
    return pd.DataFrame(rows)


# ── Deliverable 2: Column_Lineage ─────────────────────────────────────────────

def build_column_lineage_df(
    source_result: SourceAnalysisResult,
    kpi_result: KpiAnalysisResult,
    config: RationalizationConfig,
) -> pd.DataFrame:
    """One row per (master column × original column source entry).

    Columns:
        Master Source Column, Source Workbook, Source Tab,
        Original Column Name, Status,
        Used By KPI Count, Used By KPI Names
    """
    # Build canonical → KPI label list
    canonical_to_kpi_labels: dict[str, list[str]] = defaultdict(list)
    for dep in kpi_result.dependencies:
        for canonical in dep.canonical_source_columns:
            if dep.kpi_label and dep.kpi_label not in canonical_to_kpi_labels[canonical]:
                canonical_to_kpi_labels[canonical].append(dep.kpi_label)

    # Build canonical → KPI dep count
    canonical_usage: dict[str, int] = defaultdict(int)
    for dep in kpi_result.dependencies:
        for canonical in dep.canonical_source_columns:
            canonical_usage[canonical] += 1

    # Status labels
    _STATUS = {
        "common":  "Merged into Master Column",
        "similar": "Kept Separate (Similar Name)",
        "unique":  "Retained as Unique Column",
    }

    rows = []
    excl_set = set(source_result.excluded_columns)

    for profile in sorted(source_result.column_profiles, key=lambda p: p.canonical_name):
        status = _STATUS.get(profile.match_class, profile.match_class.title())
        kpi_labels = canonical_to_kpi_labels.get(profile.canonical_name, [])
        kpi_count = canonical_usage.get(profile.canonical_name, 0)

        for wb, tab, orig in profile.source_entries:
            # Override status for excluded columns
            effective_status = status
            if (wb, tab, orig) in excl_set:
                effective_status = "Excluded (Not KPI-Referenced)"

            rows.append({
                "Master Source Column": profile.canonical_name,
                "Source Workbook":      wb,
                "Source Tab":           tab,
                "Original Column Name": orig,
                "Status":               effective_status,
                "Used By KPI Count":    kpi_count,
                "Used By KPI Names":    "; ".join(kpi_labels),
            })

    if not rows:
        return pd.DataFrame(columns=[
            "Master Source Column", "Source Workbook", "Source Tab",
            "Original Column Name", "Status",
            "Used By KPI Count", "Used By KPI Names",
        ])
    return pd.DataFrame(rows)


# ── Deliverable 3: KPI_Dependencies ──────────────────────────────────────────

def build_kpi_dependencies_df(
    source_result: SourceAnalysisResult,
    kpi_result: KpiAnalysisResult,
    config: RationalizationConfig,
) -> pd.DataFrame:
    """One row per KPI formula — compact dependency map.

    Columns:
        KPI Workbook, KPI Tab, KPI Label,
        Source Columns Used, Master Columns Used,
        Aggregate Function, Formula Type, Confidence
    """
    reverse = _build_reverse_mapping(source_result)
    rows = []

    for dep in kpi_result.dependencies:
        # Original columns for this workbook
        wb_cfg = config.config_for(dep.workbook_name)
        src_tab = wb_cfg.source_tab if wb_cfg else ""
        orig_cols: list[str] = []
        for canonical in dep.canonical_source_columns:
            entries = [
                orig for wb, tab, orig in reverse.get(canonical, [])
                if wb == dep.workbook_name and tab == src_tab
            ]
            orig_cols.extend(entries if entries else [canonical])

        formula_type = "CROSS_SHEET" if not dep.refs_source_tab_only else "SINGLE_SOURCE"

        rows.append({
            "KPI Workbook":        dep.workbook_name,
            "KPI Tab":             dep.kpi_tab,
            "KPI Label":           dep.kpi_label,
            "Source Columns Used": "; ".join(dict.fromkeys(orig_cols)),
            "Master Columns Used": "; ".join(dep.canonical_source_columns),
            "Aggregate Function":  dep.aggregate_function,
            "Formula Type":        formula_type,
            "Confidence":          _kpi_confidence(dep),
        })

    if not rows:
        return pd.DataFrame(columns=[
            "KPI Workbook", "KPI Tab", "KPI Label",
            "Source Columns Used", "Master Columns Used",
            "Aggregate Function", "Formula Type", "Confidence",
        ])
    return pd.DataFrame(rows)


# ── Workbook assembly ─────────────────────────────────────────────────────────

def build_lineage_workbook(
    source_result: SourceAnalysisResult,
    kpi_result: KpiAnalysisResult,
    config: RationalizationConfig,
) -> bytes:
    """Build a standalone Data_Lineage.xlsx workbook and return bytes.

    Sheets:
        01_Data_Lineage      — KPI → source column traceability
        02_Column_Lineage    — Master column provenance
        03_KPI_Dependencies  — Compact dependency map
        Legend               — Confidence colour key
    """
    wb = Workbook()
    wb.remove(wb.active)

    lineage_df = build_data_lineage_df(source_result, kpi_result, config)
    col_df = build_column_lineage_df(source_result, kpi_result, config)
    dep_df = build_kpi_dependencies_df(source_result, kpi_result, config)

    _write_lineage_sheet(wb, "01_Data_Lineage", lineage_df,
                         "Data Lineage — KPI to Source Column Traceability",
                         conf_col="Confidence Score")
    _write_lineage_sheet(wb, "02_Column_Lineage", col_df,
                         "Column Lineage — Master Column Provenance")
    _write_lineage_sheet(wb, "03_KPI_Dependencies", dep_df,
                         "KPI Dependencies — Compact Dependency Map",
                         conf_col="Confidence")

    _write_legend(wb)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ── Sheet writers ─────────────────────────────────────────────────────────────

def _write_lineage_sheet(
    wb: Workbook,
    sheet_name: str,
    df: pd.DataFrame,
    title: str,
    conf_col: str | None = None,
) -> None:
    ws = wb.create_sheet(sheet_name)
    if df.empty:
        ws.cell(row=1, column=1, value=title).font = Font(bold=True, size=13)
        ws.cell(row=3, column=1, value="No data available.").font = Font(italic=True)
        return

    # Title row
    ws.cell(row=1, column=1, value=title).font = Font(bold=True, size=13)

    # Header row
    headers = list(df.columns)
    for c_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=3, column=c_idx, value=h)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(wrap_text=True)
    ws.row_dimensions[3].height = 20

    # Data rows
    conf_col_idx = headers.index(conf_col) + 1 if conf_col and conf_col in headers else None
    for r_idx, row in enumerate(df.itertuples(index=False), start=4):
        fill = _ALT_FILL if (r_idx % 2 == 0) else None
        for c_idx, val in enumerate(row, start=1):
            cell = ws.cell(row=r_idx, column=c_idx, value=val)
            cell.alignment = Alignment(wrap_text=True)
            if fill:
                cell.fill = fill
        # Colour confidence column
        if conf_col_idx:
            conf_val = str(df.iloc[r_idx - 4, conf_col_idx - 1])
            try:
                pct = float(conf_val.strip("%")) / 100
            except ValueError:
                pct = 1.0
            conf_cell = ws.cell(row=r_idx, column=conf_col_idx)
            if pct >= 0.95:
                conf_cell.fill = _HIGH_FILL
            elif pct >= 0.70:
                conf_cell.fill = _MED_FILL
            else:
                conf_cell.fill = _LOW_FILL

    # Column widths
    for col in ws.columns:
        max_len = max((len(str(c.value or "")) for c in col), default=10)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 60)


def _write_legend(wb: Workbook) -> None:
    ws = wb.create_sheet("Legend")
    ws.cell(row=1, column=1, value="Data Lineage — Legend").font = Font(bold=True, size=13)
    ws.cell(row=3, column=1, value="Confidence Colour Key").font = Font(bold=True)
    entries = [
        ("100% — All source columns resolved; formula references source tab only.", _HIGH_FILL),
        ("75%  — Formula references sheet(s) other than the configured source tab.", _MED_FILL),
        ("40%  — No source columns could be resolved.", _LOW_FILL),
    ]
    for row_idx, (label, fill) in enumerate(entries, start=4):
        cell = ws.cell(row=row_idx, column=1, value=label)
        cell.fill = fill

    ws.cell(row=8, column=1, value="Column Status Values").font = Font(bold=True)
    statuses = [
        ("Merged into Master Column",    "Exact normalized name match in ≥2 workbooks — auto-merged."),
        ("Kept Separate (Similar Name)", "Fuzzy-similar name but treated as distinct business field."),
        ("Retained as Unique Column",    "Appears in only one workbook — kept as-is."),
        ("Excluded (Not KPI-Referenced)","Column present in source but not used by any KPI formula."),
    ]
    for row_idx, (status, desc) in enumerate(statuses, start=9):
        ws.cell(row=row_idx, column=1, value=status).font = Font(bold=True)
        ws.cell(row=row_idx, column=2, value=desc)

    ws.column_dimensions["A"].width = 50
    ws.column_dimensions["B"].width = 65
