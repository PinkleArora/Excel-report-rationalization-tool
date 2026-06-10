"""Tests for src.ingestion.loader."""

from __future__ import annotations

import textwrap
from pathlib import Path

import pandas as pd
import pytest

from src.ingestion.loader import (
    SUPPORTED_EXTENSIONS,
    WorkbookBundle,
    load_many,
    load_workbook,
)


# ---------------------------------------------------------------------------
# Fixtures
# ---------------------------------------------------------------------------

@pytest.fixture
def simple_xlsx(tmp_path: Path) -> Path:
    """A two-sheet .xlsx file."""
    path = tmp_path / "test_wb.xlsx"
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        pd.DataFrame({"a": [1, 2, 3], "b": ["x", "y", "z"]}).to_excel(writer, sheet_name="Sheet1", index=False)
        pd.DataFrame({"c": [10.0, 20.0], "d": [True, False]}).to_excel(writer, sheet_name="Sheet2", index=False)
    return path


@pytest.fixture
def simple_csv(tmp_path: Path) -> Path:
    path = tmp_path / "test.csv"
    path.write_text("name,value\nalpha,1\nbeta,2\n")
    return path


@pytest.fixture
def formula_xlsx(tmp_path: Path) -> Path:
    """Workbook that contains a formula cell."""
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Formulas"
    ws["A1"] = "value"
    ws["B1"] = "doubled"
    ws["A2"] = 5
    ws["B2"] = "=A2*2"
    ws["A3"] = 10
    ws["B3"] = "=A3*2"
    path = tmp_path / "formulas.xlsx"
    wb.save(path)
    return path


# ---------------------------------------------------------------------------
# WorkbookBundle dataclass
# ---------------------------------------------------------------------------

class TestWorkbookBundle:
    def test_sheet_names(self):
        df = pd.DataFrame()
        bundle = WorkbookBundle(source_path=Path("x.xlsx"), sheets={"A": df, "B": df})
        assert bundle.sheet_names == ["A", "B"]

    def test_file_name(self):
        bundle = WorkbookBundle(source_path=Path("/some/path/report.xlsx"))
        assert bundle.file_name == "report.xlsx"

    def test_empty_sheets(self):
        bundle = WorkbookBundle(source_path=Path("x.xlsx"))
        assert bundle.sheet_names == []


# ---------------------------------------------------------------------------
# load_workbook
# ---------------------------------------------------------------------------

class TestLoadWorkbook:
    def test_loads_xlsx_sheet_count(self, simple_xlsx):
        bundle = load_workbook(simple_xlsx)
        assert len(bundle.sheets) == 2

    def test_loads_xlsx_sheet_names(self, simple_xlsx):
        bundle = load_workbook(simple_xlsx)
        assert set(bundle.sheet_names) == {"Sheet1", "Sheet2"}

    def test_loads_xlsx_dataframe_shape(self, simple_xlsx):
        bundle = load_workbook(simple_xlsx)
        assert bundle.sheets["Sheet1"].shape == (3, 2)

    def test_loads_csv_single_sheet(self, simple_csv):
        bundle = load_workbook(simple_csv)
        assert len(bundle.sheets) == 1
        assert "test" in bundle.sheet_names[0]

    def test_loads_csv_data(self, simple_csv):
        bundle = load_workbook(simple_csv)
        df = list(bundle.sheets.values())[0]
        assert list(df.columns) == ["name", "value"]
        assert len(df) == 2

    def test_missing_file_raises_file_not_found(self, tmp_path):
        with pytest.raises(FileNotFoundError):
            load_workbook(tmp_path / "ghost.xlsx")

    def test_unsupported_extension_raises_value_error(self, tmp_path):
        bad = tmp_path / "file.docx"
        bad.write_text("dummy")
        with pytest.raises(ValueError, match="Unsupported"):
            load_workbook(bad)

    def test_source_path_stored(self, simple_xlsx):
        bundle = load_workbook(simple_xlsx)
        assert bundle.source_path == simple_xlsx

    def test_raw_wb_available_for_xlsx(self, simple_xlsx):
        bundle = load_workbook(simple_xlsx)
        assert bundle._raw_wb is not None

    def test_supported_extensions_set(self):
        assert ".xlsx" in SUPPORTED_EXTENSIONS
        assert ".csv" in SUPPORTED_EXTENSIONS
        assert ".xls" in SUPPORTED_EXTENSIONS


# ---------------------------------------------------------------------------
# load_many
# ---------------------------------------------------------------------------

class TestLoadMany:
    def test_loads_multiple(self, simple_xlsx, simple_csv):
        bundles = load_many([simple_xlsx, simple_csv])
        assert len(bundles) == 2

    def test_skips_missing_files(self, simple_xlsx, tmp_path):
        bundles = load_many([simple_xlsx, tmp_path / "missing.xlsx"])
        assert len(bundles) == 1

    def test_empty_list(self):
        assert load_many([]) == []
