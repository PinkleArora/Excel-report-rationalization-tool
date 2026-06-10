"""Tests for src.profiling — metadata extraction and summary generation."""

from __future__ import annotations

from pathlib import Path

import pandas as pd
import pytest

from src.ingestion.loader import WorkbookBundle, load_workbook
from src.profiling.metadata import ColumnMetadata, SheetMetadata, WorkbookMetadata
from src.profiling.profiler import (
    build_summary_dataframes,
    profile_many,
    profile_sheet,
    profile_workbook,
)


# ---------------------------------------------------------------------------
# Fixtures
# ---------------------------------------------------------------------------

@pytest.fixture
def mixed_df() -> pd.DataFrame:
    return pd.DataFrame({
        "int_col": [1, 2, 3, 4, 5],
        "float_col": [1.1, 2.2, None, 4.4, 5.5],
        "str_col": ["a", "b", "c", "d", "e"],
        "bool_col": [True, False, True, False, True],
        "empty_col": [None, None, None, None, None],
    })


@pytest.fixture
def bundle_with_formula(tmp_path: Path) -> WorkbookBundle:
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws["A1"] = "x"
    ws["B1"] = "y"
    ws["A2"] = 3
    ws["B2"] = "=A2+1"
    ws["A3"] = 7
    ws["B3"] = "=A3+1"
    path = tmp_path / "formula_wb.xlsx"
    wb.save(path)
    return load_workbook(path)


@pytest.fixture
def simple_bundle(tmp_path: Path) -> WorkbookBundle:
    path = tmp_path / "simple.xlsx"
    df1 = pd.DataFrame({"name": ["Alice", "Bob"], "score": [90, 85]})
    df2 = pd.DataFrame({"dept": ["Eng", "HR"], "headcount": [10, 5]})
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        df1.to_excel(writer, sheet_name="Scores", index=False)
        df2.to_excel(writer, sheet_name="Headcount", index=False)
    return load_workbook(path)


# ---------------------------------------------------------------------------
# profile_sheet
# ---------------------------------------------------------------------------

class TestProfileSheet:
    def test_returns_sheet_metadata(self, mixed_df):
        result = profile_sheet(mixed_df, sheet_name="Test", workbook_name="wb.xlsx")
        assert isinstance(result, SheetMetadata)

    def test_row_count(self, mixed_df):
        sm = profile_sheet(mixed_df, sheet_name="T", workbook_name="wb.xlsx")
        assert sm.row_count == 5

    def test_col_count(self, mixed_df):
        sm = profile_sheet(mixed_df, sheet_name="T", workbook_name="wb.xlsx")
        assert sm.col_count == 5

    def test_empty_col_count(self, mixed_df):
        sm = profile_sheet(mixed_df, sheet_name="T", workbook_name="wb.xlsx")
        assert sm.empty_col_count == 1  # empty_col is all-None

    def test_column_names_preserved(self, mixed_df):
        sm = profile_sheet(mixed_df, sheet_name="T", workbook_name="wb.xlsx")
        assert sm.column_names == list(mixed_df.columns)

    def test_column_profiles_count(self, mixed_df):
        sm = profile_sheet(mixed_df, sheet_name="T", workbook_name="wb.xlsx")
        assert len(sm.columns) == 5

    def test_null_counts(self, mixed_df):
        sm = profile_sheet(mixed_df, sheet_name="T", workbook_name="wb.xlsx")
        float_col = next(c for c in sm.columns if c.column_name == "float_col")
        assert float_col.null_count == 1
        assert float_col.non_null_count == 4

    def test_null_pct(self, mixed_df):
        sm = profile_sheet(mixed_df, sheet_name="T", workbook_name="wb.xlsx")
        float_col = next(c for c in sm.columns if c.column_name == "float_col")
        assert abs(float_col.null_pct - 20.0) < 0.01

    def test_is_empty_flag(self, mixed_df):
        sm = profile_sheet(mixed_df, sheet_name="T", workbook_name="wb.xlsx")
        empty_col = next(c for c in sm.columns if c.column_name == "empty_col")
        assert empty_col.is_empty is True
        int_col = next(c for c in sm.columns if c.column_name == "int_col")
        assert int_col.is_empty is False

    def test_sample_values_length(self, mixed_df):
        sm = profile_sheet(mixed_df, sheet_name="T", workbook_name="wb.xlsx")
        int_col = next(c for c in sm.columns if c.column_name == "int_col")
        assert len(int_col.sample_values) == 5

    def test_sample_values_empty_col(self, mixed_df):
        sm = profile_sheet(mixed_df, sheet_name="T", workbook_name="wb.xlsx")
        empty_col = next(c for c in sm.columns if c.column_name == "empty_col")
        assert empty_col.sample_values == []

    def test_semantic_type_integer(self, mixed_df):
        sm = profile_sheet(mixed_df, sheet_name="T", workbook_name="wb.xlsx")
        int_col = next(c for c in sm.columns if c.column_name == "int_col")
        assert int_col.semantic_type == "integer"

    def test_semantic_type_float(self, mixed_df):
        sm = profile_sheet(mixed_df, sheet_name="T", workbook_name="wb.xlsx")
        col = next(c for c in sm.columns if c.column_name == "float_col")
        assert col.semantic_type == "float"

    def test_semantic_type_text(self, mixed_df):
        sm = profile_sheet(mixed_df, sheet_name="T", workbook_name="wb.xlsx")
        col = next(c for c in sm.columns if c.column_name == "str_col")
        assert col.semantic_type == "text"

    def test_semantic_type_boolean(self, mixed_df):
        sm = profile_sheet(mixed_df, sheet_name="T", workbook_name="wb.xlsx")
        col = next(c for c in sm.columns if c.column_name == "bool_col")
        assert col.semantic_type == "boolean"

    def test_semantic_type_empty(self, mixed_df):
        sm = profile_sheet(mixed_df, sheet_name="T", workbook_name="wb.xlsx")
        col = next(c for c in sm.columns if c.column_name == "empty_col")
        assert col.semantic_type == "empty"

    def test_mean_computed_for_numeric(self, mixed_df):
        sm = profile_sheet(mixed_df, sheet_name="T", workbook_name="wb.xlsx")
        int_col = next(c for c in sm.columns if c.column_name == "int_col")
        assert int_col.mean_value == pytest.approx(3.0)

    def test_min_max_computed_for_numeric(self, mixed_df):
        sm = profile_sheet(mixed_df, sheet_name="T", workbook_name="wb.xlsx")
        int_col = next(c for c in sm.columns if c.column_name == "int_col")
        assert int_col.min_value == 1
        assert int_col.max_value == 5

    def test_unique_count(self, mixed_df):
        sm = profile_sheet(mixed_df, sheet_name="T", workbook_name="wb.xlsx")
        str_col = next(c for c in sm.columns if c.column_name == "str_col")
        assert str_col.unique_count == 5

    def test_formula_count_via_raw_ws(self, bundle_with_formula):
        raw_ws = bundle_with_formula._raw_wb["Data"]
        df = bundle_with_formula.sheets["Data"]
        sm = profile_sheet(df, sheet_name="Data", workbook_name="formula_wb.xlsx", raw_ws=raw_ws)
        assert sm.formula_cell_count == 2

    def test_workbook_name_set(self, mixed_df):
        sm = profile_sheet(mixed_df, sheet_name="T", workbook_name="my_wb.xlsx")
        assert sm.workbook_name == "my_wb.xlsx"
        assert all(c.workbook_name == "my_wb.xlsx" for c in sm.columns)

    def test_empty_dataframe(self):
        sm = profile_sheet(pd.DataFrame(), sheet_name="Empty", workbook_name="wb.xlsx")
        assert sm.row_count == 0
        assert sm.col_count == 0
        assert sm.columns == []

    def test_single_row_dataframe(self):
        df = pd.DataFrame({"x": [42]})
        sm = profile_sheet(df, sheet_name="S", workbook_name="wb.xlsx")
        assert sm.row_count == 1
        assert sm.columns[0].sample_values == [42]


# ---------------------------------------------------------------------------
# profile_workbook
# ---------------------------------------------------------------------------

class TestProfileWorkbook:
    def test_returns_workbook_metadata(self, simple_bundle):
        wm = profile_workbook(simple_bundle)
        assert isinstance(wm, WorkbookMetadata)

    def test_sheet_count(self, simple_bundle):
        wm = profile_workbook(simple_bundle)
        assert wm.sheet_count == 2

    def test_workbook_name(self, simple_bundle):
        wm = profile_workbook(simple_bundle)
        assert wm.workbook_name == "simple.xlsx"

    def test_total_rows(self, simple_bundle):
        wm = profile_workbook(simple_bundle)
        assert wm.total_rows == 4  # 2 rows in each sheet

    def test_total_cols(self, simple_bundle):
        wm = profile_workbook(simple_bundle)
        assert wm.total_cols == 4  # 2 cols in each sheet

    def test_sheets_list_populated(self, simple_bundle):
        wm = profile_workbook(simple_bundle)
        assert len(wm.sheets) == 2
        sheet_names = {s.sheet_name for s in wm.sheets}
        assert sheet_names == {"Scores", "Headcount"}

    def test_file_extension(self, simple_bundle):
        wm = profile_workbook(simple_bundle)
        assert wm.file_extension == ".xlsx"


# ---------------------------------------------------------------------------
# profile_many
# ---------------------------------------------------------------------------

class TestProfileMany:
    def test_profiles_multiple_bundles(self, simple_bundle):
        results = profile_many([simple_bundle, simple_bundle])
        assert len(results) == 2

    def test_skips_bad_bundle(self, simple_bundle, monkeypatch):
        def boom(b):
            raise RuntimeError("oops")
        monkeypatch.setattr("src.profiling.profiler.profile_workbook", boom)
        results = profile_many([simple_bundle])
        assert results == []

    def test_empty_list(self):
        assert profile_many([]) == []


# ---------------------------------------------------------------------------
# build_summary_dataframes
# ---------------------------------------------------------------------------

class TestBuildSummaryDataframes:
    @pytest.fixture
    def three_dfs(self, simple_bundle):
        wm = profile_workbook(simple_bundle)
        return build_summary_dataframes([wm])

    def test_returns_three_dataframes(self, three_dfs):
        assert len(three_dfs) == 3
        for df in three_dfs:
            assert isinstance(df, pd.DataFrame)

    def test_workbook_summary_columns(self, three_dfs):
        wb_df, _, _ = three_dfs
        expected = {"workbook_name", "sheet_count", "total_rows", "total_cols"}
        assert expected.issubset(set(wb_df.columns))

    def test_workbook_summary_row_count(self, three_dfs):
        wb_df, _, _ = three_dfs
        assert len(wb_df) == 1

    def test_sheet_summary_row_count(self, three_dfs):
        _, sheet_df, _ = three_dfs
        assert len(sheet_df) == 2

    def test_column_summary_has_all_cols(self, three_dfs):
        _, _, col_df = three_dfs
        # simple_bundle has 2 sheets × 2 columns = 4 column rows
        assert len(col_df) == 4

    def test_column_summary_has_required_fields(self, three_dfs):
        _, _, col_df = three_dfs
        required = {"workbook_name", "sheet_name", "column_name", "inferred_dtype",
                    "semantic_type", "null_pct", "sample_values"}
        assert required.issubset(set(col_df.columns))

    def test_empty_input_returns_empty_dfs(self):
        wb_df, sheet_df, col_df = build_summary_dataframes([])
        assert len(wb_df) == 0
        assert len(sheet_df) == 0
        assert len(col_df) == 0
